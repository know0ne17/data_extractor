from flask import Flask,redirect,url_for,render_template,request,send_file
from flask_uploads import UploadSet,configure_uploads,DOCUMENTS
# from flask_reuploaded import UploadSet,configure_uploads,DOCUMENTS
from werkzeug.utils import secure_filename
import os
from docx import Document
from pdfplumber import open as open_pdf
import re
import io
from openpyxl import Workbook
# import jsonify

app=Flask(__name__)


# config. the upload dir
app.config['UPLOADS_DEFAULT_DEST'] = 'uploads'
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'docx'}
documents = UploadSet('documents', DOCUMENTS)
configure_uploads(app,documents)


@app.route('/upload', methods=['POST'])
def upload_file():
    if 'files' not in request.files:
        return 'No file part'
    

    uploaded_files= request.files.getlist('files')
    # file = request.files['files']
    # if file.filename == '':
    #     return 'No selected file'

    wb=Workbook()
    ws=wb.active
    ws.title = "Information"
    headers =["Contact Number","Email","Text"]
    ws.append(headers)

    for file in uploaded_files:
        filename = secure_filename(file.filename)
        file.save(os.path.join(app.config['UPLOADS_DEFAULT_DEST'], filename))

        # Extract data from the resume
        if filename.endswith('.docx'):
            doc = Document(os.path.join(app.config['UPLOADS_DEFAULT_DEST'], filename))
            text = ' '.join([paragraph.text for paragraph in doc.paragraphs])
        elif filename.endswith('.pdf'):
            text = ''
            with open_pdf(os.path.join(app.config['UPLOADS_DEFAULT_DEST'], filename)) as pdf:
                for page in pdf.pages:
                    text += page.extract_text()
        else:
            return 'Unsupported file format'

        contact = re.search(r'\s*(\d+\W\d+)', text)
        if contact:
            contact = contact.group(1)
        else:
            contact = ''
        email = re.search(r'\s*(\w+@\w+\.\w+)', text)
        if email:
            email = email.group(1)
        else:
            email = ''

        ws.append([contact, email, text])

        # Remove the uploaded file
        os.remove(os.path.join(app.config['UPLOADS_DEFAULT_DEST'], filename))
    # save excel
    excel_file_path = "contact_information.xlsx"
    wb.save(excel_file_path)

    # return 'finished'
    return send_file(excel_file_path,as_attachment=True)

@app.route('/',)
def home():
    return render_template("index.html")

if __name__ == '__main__':
    app.run()